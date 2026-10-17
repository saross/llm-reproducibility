#!/usr/bin/env python3
"""Compare reproduced values with published values: herskind-riede-2024, attempt 04.

Purpose
-------
Scores the locked targets T01-T18 of the approved plan against the values
transcribed into ``published-values/`` (before the first run) and the
deposited ``Herskind&Riede_S3.xlsx``, reading reproduced values only from the
sealed outputs of the final lane run (``outputs/<run>/``). It never runs the
authors' code and never recomputes an analysis: it compares numbers and
strings, and writes one value-level detail file per target under
``comparisons/`` plus ``comparisons/summary.json``.

Usage
-----
    venv/bin/python -I comparisons/compare.py <attempt dir> run-04

Requires openpyxl (to read S3.xlsx).
"""
from __future__ import annotations

import csv
import hashlib
import json
import math
import sys
from decimal import ROUND_HALF_UP, Decimal
from pathlib import Path

import openpyxl

ATTEMPT = Path(sys.argv[1])
RUN = sys.argv[2]
OUT = ATTEMPT / "outputs" / RUN
PUB = ATTEMPT / "published-values"
CMP = ATTEMPT / "comparisons"
SUMMARY: dict[str, dict] = {}
LISTED = {f["path"]: f["sha256"] for f in json.loads(
    (ATTEMPT / "lane-records" / RUN / "outputs.json").read_text(encoding="utf-8"))["files"]}


def read_csv(path: Path) -> list[dict]:
    """Read a CSV file into a list of dicts (UTF-8)."""
    with path.open(encoding="utf-8", newline="") as handle:
        return list(csv.DictReader(handle))


def write_csv(name: str, rows: list[dict]) -> str:
    """Write a detail CSV under comparisons/ and return its relative path."""
    path = CMP / name
    with path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(rows[0]))
        writer.writeheader()
        writer.writerows(rows)
    return f"comparisons/{name}"


def cite(rel: str) -> dict:
    """A citation of a sealed run output, checked against outputs.json."""
    digest = hashlib.sha256((OUT / rel).read_bytes()).hexdigest()
    listed = LISTED.get(rel)
    if listed is not None and listed != digest:
        raise SystemExit(f"{rel} changed after the run sealed it")
    return {"run": RUN, "path": rel, "sha256": digest}


def sig15(x: float) -> str:
    """Format a number to 15 significant digits (fwrite / Excel storage)."""
    return f"{float(x):.15g}"


def agree15(a: float, b: float) -> bool:
    """True when a and b agree to 15 significant digits.

    The test is numeric: the difference is at most half a unit in the 15th
    significant digit. A string comparison of two 15-digit renderings fails
    on rounding ties (fwrite and printf round a trailing ...5 differently).
    """
    a, b = float(a), float(b)
    scale = max(abs(a), abs(b))
    if scale == 0.0:
        return a == b
    half_unit = 0.5 * 10 ** (math.floor(math.log10(scale)) - 14)
    return abs(a - b) <= half_unit * (1 + 1e-9)


def round2(x: float) -> str:
    """Round half-up to 2 decimal places, from the shortest repr of x."""
    return str(Decimal(repr(float(x))).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP))


def summarise(tid: str, compared: int, matched: int, detail: str, outputs: list[dict],
              max_abs: float | None = None, **extra) -> None:
    """Record a target's counts for the comparison record."""
    SUMMARY[tid] = {"values_compared": compared, "values_matched": matched,
                    "evidence": detail, "outputs": outputs, "max_abs_difference": max_abs,
                    **extra}


# ---------------------------------------------------------------- T02 Fig. 3
pub3 = read_csv(PUB / "fig3-published.csv")
rep3 = {(r["panel"], int(r["rank"])): r for r in read_csv(OUT / "files/captured/fig3-reproduced.csv")}
rows, matched, comp_ok, order_ok = [], 0, 0, 0
for p in pub3:
    r = rep3.get((p["panel"], int(p["rank"])), {})
    label_ok = r.get("pmi_label") == p["pmi_label"]
    total_ok = r and int(r["total"]) == int(p["total"])
    ident_ok = r.get("ngram") == p["ngram"]
    comp = all(r and int(r[k]) == int(p[k]) for k in
               ("ertebolle", "kongemose", "kongemose_ertebolle", "maglemose"))
    matched += label_ok + bool(total_ok)
    comp_ok += comp
    order_ok += ident_ok
    rows.append({"panel": p["panel"], "rank": p["rank"], "published_ngram": p["ngram"],
                 "reproduced_ngram": r.get("ngram"), "ngram_match": ident_ok,
                 "published_pmi_label": p["pmi_label"],
                 "reproduced_pmi_label": r.get("pmi_label"),
                 "reproduced_pmi_full": r.get("pmi_full"), "pmi_label_match": label_ok,
                 "published_total": p["total"], "reproduced_total": r.get("total"),
                 "total_match": bool(total_ok),
                 "published_composition_E_K_KE_M": "/".join(p[k] for k in (
                     "ertebolle", "kongemose", "kongemose_ertebolle", "maglemose")),
                 "reproduced_composition_E_K_KE_M": "/".join(str(r.get(k)) for k in (
                     "ertebolle", "kongemose", "kongemose_ertebolle", "maglemose")),
                 "composition_match": comp})
summarise("T02", 2 * len(pub3), matched, write_csv("t02-fig3-bars.csv", rows),
          [cite("files/captured/fig3-reproduced.csv")] +
          [cite(f"files/deposit/{n}.png") for n in (
              "bigramsBYfrequency", "bigramsBYpmi", "trigramsBYfrequency", "trigramsBYpmi",
              "quadrigramsBYfrequency", "quadrigramsBYpmi")],
          bars=len(pub3), identity_order_matched=order_ok, composition_matched=comp_ok)

# ------------------------------------------------------------- T05 Table 1
pub1 = read_csv(PUB / "table1-published.csv")
rep1 = {(r["level"], int(r["frequency"])): int(r["n"])
        for r in read_csv(OUT / "files/captured/table1-reproduced.csv")}
rows, matched = [], 0
for p in pub1:
    key = (p["level"], int(p["frequency"]))
    got = rep1.get(key)
    freq_ok = got is not None          # the frequency row exists in the reproduction
    n_ok = got == int(p["n"])
    matched += freq_ok + n_ok
    rows.append({"level": p["level"], "published_frequency": p["frequency"],
                 "frequency_present": freq_ok, "published_n": p["n"], "reproduced_n": got,
                 "n_match": n_ok})
extra_rows = sorted(set(rep1) - {(p["level"], int(p["frequency"])) for p in pub1})
tot_pub = {r["level"]: int(r["n_types"]) for r in read_csv(PUB / "table1-totals-published.csv")}
tot_rep = {r["level"]: int(r["n_types"])
           for r in read_csv(OUT / "files/captured/table1-totals-reproduced.csv")}
for level, n in tot_pub.items():
    ok = tot_rep.get(level) == n
    matched += ok
    rows.append({"level": level, "published_frequency": "header n (types)",
                 "frequency_present": "", "published_n": n, "reproduced_n": tot_rep.get(level),
                 "n_match": ok})
summarise("T05", 2 * len(pub1) + len(tot_pub), matched, write_csv("t05-table1.csv", rows),
          [cite("files/captured/table1-reproduced.csv"),
           cite("files/captured/table1-totals-reproduced.csv"), cite("stdout.log")],
          reproduced_rows_not_in_table=[f"{a}:{b}" for a, b in extra_rows])

# ------------------------------------------------------------- T06 Table 2
pub2 = read_csv(PUB / "table2-published.csv")
rep2 = {r["statistic"]: r for r in read_csv(OUT / "files/captured/table2-reproduced.csv")}
stat_map = {"maximum": "Max.", "third_quartile": "3rd Qu.", "mean": "Mean",
            "median": "Median", "first_quartile": "1st Qu.", "minimum": "Min."}
rows, matched, maxdiff = [], 0, 0.0
for p in pub2:
    for col in ("maglemose", "kongemose", "ertebolle"):
        full = float(rep2[stat_map[p["statistic"]]][col])
        r2 = round2(full)
        # summary()'s printed value (4 significant digits), then 2 dp.
        printed = float(f"{full:.4g}")
        ok = Decimal(r2) == Decimal(p[col])
        matched += ok
        maxdiff = max(maxdiff, abs(full - float(p[col])))
        rows.append({"statistic": p["statistic"], "complex": col, "published": p[col],
                     "reproduced_full": repr(full), "reproduced_2dp": r2,
                     "summary_print_4sig": f"{printed:g}",
                     "category": "EXACT_MATCH" if ok else "MISMATCH"})
summarise("T06", len(rows), matched, write_csv("t06-table2.csv", rows),
          [cite("files/captured/table2-reproduced.csv"), cite("stdout.log")],
          max_abs=round(maxdiff, 6))

# ---------------------------------------------------------------- T08 S3
wb = openpyxl.load_workbook(ATTEMPT / "deposit/Herskind&Riede_S3.xlsx", read_only=True,
                            data_only=True)
rows, matched, order_same, keys_same = [], 0, {}, {}
for sheet, fname, n in (("bigrams", "Bigrams.csv", 2), ("trigrams", "Trigrams.csv", 3),
                        ("quadrigrams", "Quadrigrams.csv", 4)):
    s3 = [r for r in list(wb[sheet].iter_rows(values_only=True))[2:] if r[0] is not None]
    rep = read_csv(OUT / "files/deposit" / fname)
    cols = list(rep[0].keys())
    rep_by_key = {tuple(r[c] for c in cols[:n]): r for r in rep}
    s3_keys = [tuple(str(v) for v in r[:n]) for r in s3]
    order_same[sheet] = s3_keys == [tuple(r[c] for c in cols[:n]) for r in rep]
    keys_same[sheet] = set(s3_keys) == set(rep_by_key)
    for r in s3:
        key = tuple(str(v) for v in r[:n])
        got = rep_by_key.get(key)
        for i, col in enumerate(cols):
            pub_v = r[i]
            rep_v = got[col] if got else None
            if col in ("expected_freq", "PMI"):
                ok = got is not None and agree15(pub_v, rep_v)
                diff = abs(float(pub_v) - float(rep_v)) if got else None
            elif i < n:
                ok = got is not None and str(pub_v) == rep_v
                diff = None
            else:
                ok = got is not None and int(pub_v) == int(rep_v)
                diff = None
            matched += ok
            rows.append({"sheet": sheet, "key": " ".join(key), "column": col,
                         "published_s3": pub_v, "reproduced_csv": rep_v,
                         "abs_difference": diff, "match": ok})
summarise("T08", len(rows), matched, write_csv("t08-s3-cells.csv", rows),
          [cite("files/deposit/Bigrams.csv"), cite("files/deposit/Trigrams.csv"),
           cite("files/deposit/Quadrigrams.csv")],
          max_abs=max((x["abs_difference"] or 0.0) for x in rows),
          row_order_identical=order_same, row_sets_identical=keys_same)

# ----------------------------------------------------------- T03 Fig. 4
m_rows = read_csv(OUT / "files/captured/fig4-matrix-m.csv")
axis = [r[""] for r in m_rows]
pub_axis = [r for r in read_csv(PUB / "fig4-axis-and-large-boxes.csv")
            if r["item"] == "axis_order_left_to_right_and_bottom_to_top"][0]["value"].split()
axis_ok = [a.upper() for a in axis] == pub_axis
s3_bi = {frozenset((str(r[0]), str(r[1]))): r[6]
         for r in list(wb["bigrams"].iter_rows(values_only=True))[2:] if r[0] is not None}
cells, mat_matched, nonzero = [], 0, 0
for i, r in enumerate(m_rows):
    for j, b in enumerate(axis):
        if j <= i:
            continue
        v = float(r[b])
        key = frozenset((axis[i].upper(), b.upper()))
        if v != 0.0 or key in s3_bi:
            nonzero += 1
            ok = key in s3_bi and agree15(v, s3_bi[key])
            mat_matched += ok
            cells.append({"pair": " ".join(sorted(key)), "reproduced_m": repr(v),
                          "s3_pmi": s3_bi.get(key),
                          "abs_difference": abs(v - s3_bi[key]) if key in s3_bi else None,
                          "match": ok})
detail_m = write_csv("t03-matrix-vs-s3.csv", cells)
lookup = {frozenset((r["x_token"], r["y_token"])): r
          for r in read_csv(OUT / "files/captured/aid-T03-axis-pair-chronology.csv")}
boxes = {frozenset((r["x_token"], r["y_token"])): r["colour_culture"]
         for r in read_csv(PUB / "fig4-boxes-published.csv")}
norm = {"Ertebølle": "Ertebolle", "Kongemose": "Kongemose", "Maglemose": "Maglemose"}
box_rows, tp, fp, fn, colour_bad = [], 0, 0, 0, 0
plotted = set(s3_bi)
for key, rec in lookup.items():
    excl = norm.get(rec["exclusive_to"], "")
    pub = boxes.get(key, "")
    if pub or excl:
        state = ("agree" if pub and excl and pub == excl else
                 "colour differs" if pub and excl else
                 "box but not exclusive" if pub else "exclusive but no box")
        tp += state == "agree"
        colour_bad += state == "colour differs"
        fp += state == "box but not exclusive"
        fn += state == "exclusive but no box"
        box_rows.append({"pair": " ".join(sorted(key)), "plotted_obs_ge_3": key in plotted,
                         "published_box": pub, "lookup_exclusive_to": rec["exclusive_to"],
                         "lookup_n_objects": rec["n_objects"],
                         "lookup_M_K_KE_E": "/".join(rec[k] for k in (
                             "maglemose", "kongemose", "kongemose_ertebolle", "ertebolle")),
                         "state": state})
detail_b = write_csv("t03-boxes-vs-lookup.csv", box_rows)
large = [r for r in read_csv(PUB / "fig4-axis-and-large-boxes.csv")
         if r["item"].startswith("large_box")]
large_rows = []
for r in large:
    motifs = r["value"].split()
    for a_i, a in enumerate(motifs):
        for b in motifs[a_i + 1:]:
            key = frozenset((a, b))
            rec = lookup.get(key, {})
            large_rows.append({"large_box": r["item"], "pair": f"{a} {b}",
                               "s3_pmi": s3_bi.get(key),
                               "lookup_exclusive_to": rec.get("exclusive_to"),
                               "lookup_n_objects": rec.get("n_objects")})
detail_l = write_csv("t03-large-boxes.csv", large_rows)
matrix_max_abs = max((c["abs_difference"] or 0.0) for c in cells)
summarise("T03", nonzero + 1 + len(box_rows), mat_matched + axis_ok + tp,
          f"{detail_m}; {detail_b}; {detail_l}",
          [cite("files/captured/fig4-matrix-m.csv"), cite("files/deposit/heatmapPMI.png"),
           cite("files/captured/aid-T03-axis-pair-chronology.csv")],
          max_abs=matrix_max_abs,
          matrix_pairs=nonzero, matrix_matched=mat_matched,
          matrix_within_precision=sum(1 for c in cells if not c["match"] and
                                      c["abs_difference"] is not None and
                                      c["abs_difference"] < 1e-14),
          axis_order_exact=axis_ok,
          boxes_published=len(boxes), boxes_agree=tp, boxes_colour_differs=colour_bad,
          boxes_not_exclusive=fp, exclusive_without_box=fn)

# ------------------------------------------------- T01 Fig. 2 (bar panel)
fill = {"#FAD02E": "ertebolle", "#D98888": "kongemose", "#E8BB8B": "kongemose_ertebolle",
        "#B4A1C1": "maglemose"}
layer = read_csv(OUT / "files/captured/fig2-layer-data.csv")
rep_bars: dict[str, dict] = {}
y_of: dict[str, float] = {}
for r in layer:
    rep_bars.setdefault(r["category"], {k: 0 for k in fill.values()})[fill[r["fill"]]] += int(r["count"])
    y_of[r["category"]] = float(r["y"])
rep_order = sorted(rep_bars, key=lambda c: -y_of[c])
pub_bars = read_csv(PUB / "fig2-bars-published.csv")
rows, within = [], 0
order_ok = [p["category"] for p in pub_bars] == rep_order
for p in pub_bars:
    rb = rep_bars.get(p["category"], {})
    seg = {k: (int(p[k]), rb.get(k)) for k in fill.values()}
    worst = max(abs(a - (b or 0)) for a, b in seg.values())
    total_rep = sum(rb.values())
    ok = worst <= 1 and abs(int(p["total"]) - total_rep) <= 1
    within += ok
    rows.append({"rank": p["rank"], "category": p["category"],
                 "published_total_measured": p["total"], "reproduced_total": total_rep,
                 **{f"{k}_published_measured": a for k, (a, _) in seg.items()},
                 **{f"{k}_reproduced": b for k, (_, b) in seg.items()},
                 "max_segment_difference": worst, "within_figure_resolution": ok})
summarise("T01", len(rows) + 1, within + order_ok, write_csv("t01-fig2-bars.csv", rows),
          [cite("files/captured/fig2-layer-data.csv"),
           cite("files/deposit/artefact_type_chronology.png"),
           cite("files/captured/aid-T01-map.png"),
           cite("files/captured/aid-T01-map-chronology-counts.csv"),
           cite("files/captured/aid-T01-map-extent.csv"),
           cite("files/captured/aid-T01-map-unplaced-by-chronology.csv")],
          category_order_exact=order_ok, bars_within_resolution=within,
          reproduced_total_objects=sum(sum(v.values()) for v in rep_bars.values()))

# ------------------------------------------------------- T07 / T18 Table 3
pub_t3 = read_csv(PUB / "table3-published.csv")
look = {int(r["table3_number"]): r for r in
        read_csv(OUT / "files/captured/aid-T07-table3-lookup.csv")}


def norm_site(s: str) -> str:
    """Site names compared without typographic punctuation or spacing."""
    return "".join(ch for ch in s.lower() if ch.isalnum())


rows, matched = [], 0
for p in pub_t3:
    r = look[int(p["number"])]
    checks = {
        "site": norm_site(p["site"]) == norm_site(r["site"] or ""),
        "region": p["region"] == r["region"],
        "artefact_type": p["artefact_type"].lower() == f"{r['material']} {r['artefact_type']}".lower(),
        "motifs": set(p["motifs"].split()) == set((r["motifs"] or "").split()),
        "figure": p["figure_in_plonka_2003"] == r["figure"],
    }
    matched += sum(checks.values())
    rows.append({"number": p["number"], "published_site": p["site"], "s1_site": r["site"],
                 "published_region": p["region"], "s1_region": r["region"],
                 "published_type": p["artefact_type"],
                 "s1_material_type": f"{r['material']} {r['artefact_type']}",
                 "published_motifs": p["motifs"], "s1_motifs": r["motifs"],
                 "published_figure": p["figure_in_plonka_2003"], "s1_figure": r["figure"],
                 "s1_matches": r["matches"], **{f"{k}_match": v for k, v in checks.items()}})
summarise("T07", 5 * len(pub_t3), matched, write_csv("t07-table3.csv", rows),
          [cite("files/captured/aid-T07-table3-lookup.csv")])
types = [f"{look[i]['material']} {look[i]['artefact_type']}" for i in range(1, 11)]
n_axes, n_shafts = types.count("antler axe"), types.count("antler shaft")
write_csv("t18-artefact-types.csv", [{"item": "antler axes", "published": 9, "reproduced": n_axes},
                                     {"item": "antler shafts", "published": 1,
                                      "reproduced": n_shafts}])
summarise("T18", 2, (n_axes == 9) + (n_shafts == 1), "comparisons/t18-artefact-types.csv",
          [cite("files/captured/aid-T07-table3-lookup.csv")])

# ---------------------------------------------------------------- T04 Fig. 5
pub5 = read_csv(PUB / "fig5-published.csv")
rep5 = {int(r["table3_number"]): r for r in read_csv(OUT / "files/captured/aid-T04-fig5-objects.csv")}
rows, pie_ok = [], 0
for p in pub5:
    r = rep5[int(p["number"])]
    ok = set(p["pie_motifs"].split()) == set(r["cluster_motifs"].split())
    pie_ok += ok
    rows.append({"number": p["number"], "published_pie": p["pie_motifs"],
                 "s1_cluster_motifs": r["cluster_motifs"], "pie_match": ok,
                 "published_x_px": p["x_px"], "published_y_px": p["y_px"],
                 "s1_lon": r["lon"], "s1_lat": r["lat"]})


def ranks(values: list[float]) -> list[float]:
    """Average ranks (1-based) for Spearman's correlation."""
    order = sorted(range(len(values)), key=lambda i: values[i])
    out = [0.0] * len(values)
    i = 0
    while i < len(order):
        j = i
        while j + 1 < len(order) and values[order[j + 1]] == values[order[i]]:
            j += 1
        for k in range(i, j + 1):
            out[order[k]] = (i + j) / 2 + 1
        i = j + 1
    return out


def spearman(a: list[float], b: list[float]) -> float:
    """Spearman's rank correlation of two equal-length lists."""
    ra, rb = ranks(a), ranks(b)
    ma, mb = sum(ra) / len(ra), sum(rb) / len(rb)
    cov = sum((x - ma) * (y - mb) for x, y in zip(ra, rb))
    return cov / (sum((x - ma) ** 2 for x in ra) * sum((y - mb) ** 2 for y in rb)) ** 0.5


rho_x = spearman([float(r["published_x_px"]) for r in rows], [float(r["s1_lon"]) for r in rows])
rho_y = spearman([-float(r["published_y_px"]) for r in rows], [float(r["s1_lat"]) for r in rows])
summarise("T04", len(rows), pie_ok, write_csv("t04-fig5.csv", rows),
          [cite("files/captured/aid-T04-fig5-objects.csv"),
           cite("files/captured/aid-T04-fig5-objects.png")],
          spearman_x_vs_lon=round(rho_x, 3), spearman_y_vs_lat=round(rho_y, 3))

# ------------------------------------------------------ named values T09-T18
aid = lambda name: {r["item"]: r["value"] for r in read_csv(OUT / f"files/captured/{name}")}
t09 = aid("aid-T09-object-count.csv")
t10 = aid("aid-T10-max-motifs.csv")
named = []


def nv(tid: str, item: str, published, reproduced, source: str) -> None:
    """One named-value comparison row."""
    named.append({"target": tid, "item": item, "published": published,
                  "reproduced": reproduced, "match": str(published) == str(reproduced),
                  "source": source})


nv("T09", "objects (S1 rows with a number)", 483, t09["S1 Data rows with an object number"],
   "aid-T09")
nv("T10", "maximum motifs on one object", 15, t10["max ntoken over the authors' tokens"],
   "aid-T10")
nv("T10", "k", 13, t10["authors' k (max)"], "aid-T10 (authors' k <- 0:13)")
tab1 = {(r["level"], int(r["frequency"])): int(r["n"])
        for r in read_csv(OUT / "files/captured/table1-reproduced.csv")}
nv("T11", "combinations occurring 44 times", 2, tab1.get(("bigrams", 44)), "table(bigramfreqs)")
max_bigram = max(f for (lvl, f) in tab1 if lvl == "bigrams")
nv("T11", "their frequency (the maximum)", 44, max_bigram, "table(bigramfreqs)")
pct = Decimal(44 * 100) / Decimal(483)
nv("T11", "share of objects (rounded %)", "9%",
   f"{pct.quantize(Decimal('1'), rounding=ROUND_HALF_UP)}%", "44 / 483, printed inputs")
bi = read_csv(OUT / "files/deposit/Bigrams.csv")
bi_sorted = sorted(bi, key=lambda r: -int(r["observed_freq"]))
top = [(f"{r['token1']}-{r['token2']}", int(r["observed_freq"])) for r in bi_sorted[:3]]
nv("T12", "the two 44-frequency combinations", "A1-C1; C1-G1",
   "; ".join(sorted(p for p, f in top if f == 44)), "Bigrams.csv (authors' export)")
nv("T12", "third highest combination", "A1-G1", top[2][0], "Bigrams.csv")
nv("T12", "its frequency", 28, top[2][1], "Bigrams.csv")
t13 = read_csv(OUT / "files/captured/aid-T13-unigram-top.csv")
top6_tf = sorted(r["token_freqs_motif"] for r in t13[:6])
top6_fr = sorted(r["authors_freqs_motif"] for r in t13[:6])
for motif in ["A1", "A24", "B10", "C1", "G1", "H1"]:
    nv("T13", f"{motif} among the six highest unigram frequencies", True,
       motif in top6_tf and motif in top6_fr, "aid-T13 (token_freqs and authors' freqs)")
t14 = read_csv(OUT / "files/captured/aid-T14-fig3b-chronology.csv")
fig3b = [r for r in read_csv(OUT / "files/captured/fig3-reproduced.csv") if r["panel"] == "b"]


def exclusive_from_plot(r: dict) -> str:
    """Culture of a Fig. 3b bar whose objects all belong to one complex."""
    parts = {k: int(r[k]) for k in ("ertebolle", "kongemose", "kongemose_ertebolle",
                                    "maglemose")}
    nonzero = [k for k, v in parts.items() if v]
    return nonzero[0] if len(nonzero) == 1 else ""


plot_excl = {r["ngram"]: exclusive_from_plot(r) for r in fig3b}
aid_excl = {r["bigram"]: norm.get(r["exclusive_to"], "").lower() for r in t14}
agree_sources = all(plot_excl[k] == aid_excl.get(k, "") for k in plot_excl)
erte = sorted(k for k, v in plot_excl.items() if v == "ertebolle")
konge = sorted(k for k, v in plot_excl.items() if v == "kongemose")
magle = sorted(k for k, v in plot_excl.items() if v == "maglemose")
nv("T14", "Fig. 3b bigrams exclusive to Maglemose", 0, len(magle), "Fig. 3b plot data")
nv("T14", "Fig. 3b bigrams specific to Kongemose", 1, len(konge), "Fig. 3b plot data")
nv("T14", "Fig. 3b bigrams unique to Ertebolle (stated count)", 9, len(erte), "Fig. 3b plot data")
nv("T14", "the Kongemose bigram", "B10 F38", konge[0] if konge else None, "Fig. 3b plot data")
for listed in ["I5 I13", "C5 D7", "C4 C12", "C12 C13", "C4 C5", "I1 I13", "C5 C13", "I1 I5"]:
    nv("T14", f"listed Ertebolle bigram {listed}", True, listed in erte, "Fig. 3b plot data")
culture_specific = [r for r in fig3b if exclusive_from_plot(r)]
nv("T15", "highest frequency among culture-specific Fig. 3b bigrams", 5,
   max(int(r["total"]) for r in culture_specific), "Fig. 3b plot data")
t16 = read_csv(OUT / "files/captured/aid-T16-cluster-pairs.csv")
plotted_pairs = [r for r in t16 if r["authors_pmi"] not in ("", "NA")]
for motif in ["C4", "C5", "C12", "C13", "D7"]:
    in_pairs = any(motif in r["pair"].split() for r in plotted_pairs)
    nv("T16", f"{motif} in a plotted (observed >= 3) cluster pair", True, in_pairs,
       "aid-T16 / authors' bigrams")
for r in t16:
    key = frozenset(r["pair"].split())
    s3v = s3_bi.get(key)
    if r["authors_pmi"] not in ("", "NA"):
        nv("T16", f"{r['pair']} PMI agrees with S3 (15 significant digits)", True,
           s3v is not None and agree15(s3v, r["authors_pmi"]), "authors' bigrams vs S3")
    nv("T16", f"{r['pair']} co-occurs only on Ertebolle objects", "Ertebølle",
       r["exclusive_to"], "aid-T16 (S1 lookup)")
t16_rows = []
for r in t16:
    key = frozenset(r["pair"].split())
    s3v = s3_bi.get(key)
    t16_rows.append({"pair": r["pair"], "authors_observed_freq": r["authors_observed_freq"],
                     "authors_pmi": r["authors_pmi"], "s3_pmi": s3v,
                     "pmi_matches_s3": (s3v is not None and r["authors_pmi"] not in ("", "NA")
                                        and agree15(s3v, r["authors_pmi"])),
                     "s1_objects": r["n_objects"], "exclusive_to": r["exclusive_to"]})
t17 = read_csv(OUT / "files/captured/aid-T17-cluster-objects.csv")
t17_set = sorted(r["figure"] for r in t17)
pub_set = sorted(r["figure_in_plonka_2003"] for r in pub_t3)
nv("T17", "objects bearing the cluster ornament strings", 10, len(t17), "aid-T17")
named_detail = write_csv("t09-t18-named-values.csv", named)
write_csv("t16-cluster-pairs.csv", t16_rows)
write_csv("t17-cluster-objects-vs-table3.csv", [
    {"published_table3_figures": " ".join(pub_set), "lookup_figures": " ".join(t17_set),
     "identical_sets": pub_set == t17_set}])
for tid in ("T09", "T10", "T11", "T12", "T13", "T14", "T15", "T16", "T17"):
    rs = [r for r in named if r["target"] == tid]
    srcs = {"T09": ["files/captured/aid-T09-object-count.csv"],
            "T10": ["files/captured/aid-T10-max-motifs.csv"],
            "T11": ["files/captured/table1-reproduced.csv"],
            "T12": ["files/deposit/Bigrams.csv"],
            "T13": ["files/captured/aid-T13-unigram-top.csv"],
            "T14": ["files/captured/fig3-reproduced.csv",
                    "files/captured/aid-T14-fig3b-chronology.csv"],
            "T15": ["files/captured/fig3-reproduced.csv",
                    "files/captured/aid-T14-fig3b-chronology.csv"],
            "T16": ["files/captured/aid-T16-cluster-pairs.csv",
                    "files/deposit/Bigrams.csv", "files/deposit/heatmapPMI.png"],
            "T17": ["files/captured/aid-T17-cluster-objects.csv"]}[tid]
    extra = {}
    if tid == "T14":
        extra = {"plot_and_lookup_agree": agree_sources, "ertebolle_exclusive": erte,
                 "unlisted_ertebolle": sorted(set(erte) - {"I5 I13", "C5 D7", "C4 C12",
                                                           "C12 C13", "C4 C5", "I1 I13",
                                                           "C5 C13", "I1 I5"})}
    if tid == "T16":
        extra = {"pairs": t16_rows}
    if tid == "T17":
        extra = {"identical_to_table3_set": pub_set == t17_set}
    detail = named_detail + ("; comparisons/t16-cluster-pairs.csv" if tid == "T16" else "") + (
        "; comparisons/t17-cluster-objects-vs-table3.csv" if tid == "T17" else "")
    summarise(tid, len(rs), sum(r["match"] for r in rs), detail, [cite(s) for s in srcs], **extra)

(CMP / "summary.json").write_text(json.dumps(SUMMARY, indent=1, default=str) + "\n",
                                  encoding="utf-8")
for tid in sorted(SUMMARY):
    s = SUMMARY[tid]
    keys = {k: v for k, v in s.items() if k not in ("outputs", "evidence", "pairs")}
    print(tid, json.dumps(keys, default=str))
